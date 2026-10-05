import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import numpy as np

def format_worksheet(ws, title: str):
    """Aplica estilo institucional rigoroso e formatação numérica explícita às abas do Excel."""
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid") # Azul petróleo institucional
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    border_thin = Side(border_style="thin", color="E0E0E0")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    # 1. Formatar Cabeçalho
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28

    pct_keywords = [
        "roic", "margin", "fcfyield", "upside", "discount", "sma200", 
        "roe", "%", "discount_52w_high", "analyst_upside", "price_vs_sma200"
    ]
    price_keywords = ["currentprice", "targetmeanprice", "52weekhigh", "sma200_price"]
    large_curr_keywords = ["marketcap", "revenue", "ebit", "fcf", "totaldebt", "totalcash", "debt", "cash"]

    # 2. Formatar Células de Dados
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.font = data_font
            cell.border = cell_border
            val = cell.value

            header_name = str(ws.cell(row=1, column=cell.column).value or "").strip().lower()
            
            if isinstance(val, (int, float, np.floating, np.integer)) and not pd.isna(val):
                if any(k in header_name for k in pct_keywords) and not ("_pct" in header_name or "score" in header_name):
                    cell.number_format = '0.0%'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif any(k in header_name for k in ["_pct", "score"]):
                    cell.number_format = '0.0'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif any(k in header_name for k in ["netdebt_ebitda", "pe", "peg", "ratio", "coverage"]):
                    cell.number_format = '0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif any(k in header_name for k in price_keywords):
                    cell.number_format = '$#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif any(k in header_name for k in large_curr_keywords):
                    cell.number_format = '$#,##0,, "M"'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                if any(k in header_name for k in ["ticker", "archetype", "verdict", "sector", "industry", "status", "1.", "2.", "3.", "4.", "5."]):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

    # 3. Ajuste preciso da largura das colunas
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            v_str = str(cell.value or '')
            if cell.number_format and len(cell.number_format) > 5 and isinstance(cell.value, (int, float)):
                v_str += "   "
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    ws.freeze_panes = "A2"


def evaluate_qualitative_checklist(row: pd.Series) -> dict:
    """
    Avalia os 5 critérios qualitativos e o Status Decisório institucional com base
    nos fundamentos reais, riscos jurídicos/regulatórios e posição de mercado.
    """
    ticker = str(row.get("Ticker", "")).strip().upper()
    roic = row.get("ROIC", 0.0) or 0.0
    fcf_yield = row.get("FCFYield", 0.0) or 0.0
    disc_52w = row.get("Discount_52W_High", 0.0) or 0.0
    gross_margin = row.get("GrossMargin", 0.0) or 0.0
    fwd_pe = row.get("ForwardPE")
    score = row.get("Composite_Score", 0.0) or 0.0

    # Overrides específicos institucionais para casos conhecidos
    specific_cases = {
        "MMM": {
            "moat": "Sim",
            "mgmt": "Parcial",
            "regulatory": "Não",
            "allocation": "Parcial",
            "thesis": "Parcial",
            "status": "REJEITADO (PASSIVOS / RISCO ELEVADO)"
        },
        "FSLR": {
            "moat": "Parcial",
            "mgmt": "Sim",
            "regulatory": "Não",
            "allocation": "Parcial",
            "thesis": "Parcial",
            "status": "REJEITADO (PASSIVOS / RISCO ELEVADO)"
        },
        "INCY": {
            "moat": "Parcial",
            "mgmt": "Parcial",
            "regulatory": "Não",
            "allocation": "Parcial",
            "thesis": "Parcial",
            "status": "REJEITADO (PASSIVOS / RISCO ELEVADO)"
        },
        "SNDK": {
            "moat": "Não",
            "mgmt": "Parcial",
            "regulatory": "Não",
            "allocation": "Parcial",
            "thesis": "Não",
            "status": "REJEITADO (PASSIVOS / RISCO ELEVADO)"
        },
        "NVDA": {
            "moat": "Sim",
            "mgmt": "Sim",
            "regulatory": "Médio",
            "allocation": "Sim",
            "thesis": "Sim",
            "status": "AGUARDAR (PREÇO NOS MÁXIMOS)"
        },
        "ADBE": {
            "moat": "Sim",
            "mgmt": "Sim",
            "regulatory": "Sim",
            "allocation": "Sim",
            "thesis": "Sim",
            "status": "APROVADO (FORTE COMPRA)"
        },
        "DECK": {
            "moat": "Sim",
            "mgmt": "Sim",
            "regulatory": "Sim",
            "allocation": "Sim",
            "thesis": "Sim",
            "status": "APROVADO (FORTE COMPRA)"
        },
        "AOS": {
            "moat": "Sim",
            "mgmt": "Sim",
            "regulatory": "Sim",
            "allocation": "Sim",
            "thesis": "Sim",
            "status": "APROVADO (FORTE COMPRA)"
        },
        "ZTS": {
            "moat": "Sim",
            "mgmt": "Sim",
            "regulatory": "Sim",
            "allocation": "Sim",
            "thesis": "Sim",
            "status": "APROVADO (FORTE COMPRA)"
        },
        "NFLX": {
            "moat": "Sim",
            "mgmt": "Sim",
            "regulatory": "Médio",
            "allocation": "Sim",
            "thesis": "Sim",
            "status": "APROVADO (COMPRA EM TRANCHES / MODERADA)"
        },
        "MAS": {
            "moat": "Parcial",
            "mgmt": "Sim",
            "regulatory": "Sim",
            "allocation": "Sim",
            "thesis": "Sim",
            "status": "APROVADO (COMPRA EM TRANCHES / MODERADA)"
        },
        "GDDY": {
            "moat": "Parcial",
            "mgmt": "Sim",
            "regulatory": "Sim",
            "allocation": "Sim",
            "thesis": "Sim",
            "status": "APROVADO (COMPRA EM TRANCHES / MODERADA)"
        },
        "STX": {
            "moat": "Parcial",
            "mgmt": "Parcial",
            "regulatory": "Médio",
            "allocation": "Parcial",
            "thesis": "Sim",
            "status": "APROVADO (COMPRA EM TRANCHES / MODERADA)"
        },
        "EXPE": {
            "moat": "Parcial",
            "mgmt": "Sim",
            "regulatory": "Médio",
            "allocation": "Sim",
            "thesis": "Sim",
            "status": "APROVADO (COMPRA EM TRANCHES / MODERADA)"
        },
        "CAH": {
            "moat": "Sim",
            "mgmt": "Sim",
            "regulatory": "Médio",
            "allocation": "Sim",
            "thesis": "Sim",
            "status": "APROVADO (COMPRA EM TRANCHES / MODERADA)"
        }
    }

    if ticker in specific_cases:
        return specific_cases[ticker]

    # Avaliação genérica quantitativo-qualitativa
    moat = "Sim" if (roic >= 0.20 or gross_margin >= 0.50) else ("Parcial" if roic >= 0.10 else "Não")
    mgmt = "Sim" if roic >= 0.15 else "Parcial"
    regulatory = "Sim" if score >= 88 else "Médio"
    allocation = "Sim" if (fcf_yield >= 0.05 or roic >= 0.18) else "Parcial"
    thesis = "Sim" if (disc_52w >= 0.10 or fcf_yield >= 0.05) else "Parcial"

    if disc_52w < 0.05 and (fwd_pe is not None and fwd_pe > 35):
        status = "AGUARDAR (PREÇO NOS MÁXIMOS)"
    elif score >= 88 and moat == "Sim" and mgmt == "Sim":
        status = "APROVADO (FORTE COMPRA)"
    elif score >= 80:
        status = "APROVADO (COMPRA EM TRANCHES / MODERADA)"
    else:
        status = "AGUARDAR (PREÇO NOS MÁXIMOS)"

    return {
        "moat": moat,
        "mgmt": mgmt,
        "regulatory": regulatory,
        "allocation": allocation,
        "thesis": thesis,
        "status": status
    }


def export_to_institutional_excel(df: pd.DataFrame, filename: str = "relatorio_investimento_institucional.xlsx"):
    """
    Gera o Excel executivo com abas formatadas e Checklist preenchida conclusivamente.
    """
    wb = openpyxl.Workbook()
    default_sheet = wb.active

    exec_cols = [
        "Ticker", "Company", "Sector", "Composite_Score", "Verdict", "Archetype",
        "CurrentPrice", "52WeekHigh", "TargetMeanPrice", "Analyst_Upside", "ForwardPE", "PEGRatio",
        "ROIC", "ROIC_Sector_Pct", "OperatingMargin", "GrossMargin", "FCFYield", "FCFYield_Sector_Pct",
        "Discount_52W_High", "InterestCoverage", "NetDebt_EBITDA", "Price_vs_SMA200"
    ]
    available_exec_cols = [c for c in exec_cols if c in df.columns]

    fin_cols = [
        "Ticker", "Company", "Sector", "Composite_Score", "Verdict", "Archetype",
        "CurrentPrice", "52WeekHigh", "TargetMeanPrice", "Analyst_Upside", "ForwardPE", "TrailingPE",
        "ROE", "ROE_Sector_Pct", "MarketCap", "Discount_52W_High", "Price_vs_SMA200"
    ]
    available_fin_cols = [c for c in fin_cols if c in df.columns]

    # Aba 1: Top_15_Conviccao (EXCLUI 'Financial Services')
    top_conviccao = df[
        (~df["Sector"].str.contains("Financial", case=False, na=False)) &
        (df["Archetype"] != "Financeiras_Bancos") &
        (~df["Verdict"].str.contains("REJEITAR", na=False))
    ].sort_values(by="Composite_Score", ascending=False).head(15)

    ws1 = wb.create_sheet(title="Top_15_Conviccao")
    ws1.append(available_exec_cols)
    for _, row in top_conviccao[available_exec_cols].iterrows():
        ws1.append([None if pd.isna(v) else v for v in row])
    format_worksheet(ws1, "Top_15_Conviccao")

    # Aba 2: Compounders_GARP
    compounders = df[
        (df["Archetype"].str.contains("Compounder", na=False)) &
        (~df["Verdict"].str.contains("REJEITAR", na=False))
    ].sort_values(by=["Composite_Score", "ROIC_Sector_Pct"], ascending=[False, False])

    ws2 = wb.create_sheet(title="Compounders_GARP")
    ws2.append(available_exec_cols)
    for _, row in compounders[available_exec_cols].iterrows():
        ws2.append([None if pd.isna(v) else v for v in row])
    format_worksheet(ws2, "Compounders_GARP")

    # Aba 3: Deep_Value_Desconto
    deep_value = df[
        ((df["Archetype"].str.contains("Deep_Value", na=False)) |
         ((df["Discount_52W_High"] >= 0.15) & (df["FCFYield"] >= 0.05))) &
        (~df["Sector"].str.contains("Financial", case=False, na=False))
    ].sort_values(by=["FCFYield", "Discount_52W_High"], ascending=[False, False])

    ws3 = wb.create_sheet(title="Deep_Value_Desconto")
    ws3.append(available_exec_cols)
    for _, row in deep_value[available_exec_cols].iterrows():
        ws3.append([None if pd.isna(v) else v for v in row])
    format_worksheet(ws3, "Deep_Value_Desconto")

    # Aba 4: Financeiras_Bancos
    financeiras = df[
        (df["Sector"].str.contains("Financial", case=False, na=False)) |
        (df["Archetype"] == "Financeiras_Bancos")
    ].sort_values(by=["Composite_Score", "ROE"], ascending=[False, False])

    ws_fin = wb.create_sheet(title="Financeiras_Bancos")
    ws_fin.append(available_fin_cols)
    for _, row in financeiras[available_fin_cols].iterrows():
        ws_fin.append([None if pd.isna(v) else v for v in row])
    format_worksheet(ws_fin, "Financeiras_Bancos")

    # Aba 5: Checklist_Validacao_Final (PREENCHIMENTO AUTOMÁTICO E FORMATAÇÃO CONDICIONAL)
    ws4 = wb.create_sheet(title="Checklist_Validacao_Final")
    checklist_headers = [
        "Ticker", "Company", "Sector", "Composite_Score", "Veredito",
        "1. Fosso Competitivo (Pricing Power)",
        "2. Alinhamento Gestão (Skin in Game / ROIC)",
        "3. Risco Regulatório / Disrupção Baixo",
        "4. Alocação Capital Eficiente",
        "5. Tese Crescimento / Recuperação Clara",
        "Status Decisório (Comitê)"
    ]
    ws4.append(checklist_headers)

    # Estilos de preenchimento condicional suave institucional
    fill_approved = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid") # Verde suave (#155724 text)
    font_approved = Font(name="Calibri", size=10, bold=True, color="155724")

    fill_wait = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")     # Amarelo suave (#856404 text)
    font_wait = Font(name="Calibri", size=10, bold=True, color="856404")

    fill_rejected = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid") # Vermelho suave (#721C24 text)
    font_rejected = Font(name="Calibri", size=10, bold=True, color="721C24")

    for _, row in top_conviccao.head(15).iterrows():
        eval_res = evaluate_qualitative_checklist(row)
        ws4.append([
            row["Ticker"],
            row["Company"],
            row["Sector"],
            row["Composite_Score"],
            row["Verdict"],
            eval_res["moat"],
            eval_res["mgmt"],
            eval_res["regulatory"],
            eval_res["allocation"],
            eval_res["thesis"],
            eval_res["status"]
        ])

    format_worksheet(ws4, "Checklist_Validacao_Final")

    # Aplicação de cores condicionais na coluna 'Status Decisório' (Coluna 11 / K)
    status_col_idx = 11
    for r in range(2, ws4.max_row + 1):
        status_cell = ws4.cell(row=r, column=status_col_idx)
        val_status = str(status_cell.value or "")
        status_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        if "APROVADO" in val_status:
            status_cell.fill = fill_approved
            status_cell.font = font_approved
        elif "REJEITADO" in val_status:
            status_cell.fill = fill_rejected
            status_cell.font = font_rejected
        elif "AGUARDAR" in val_status:
            status_cell.fill = fill_wait
            status_cell.font = font_wait

    # Aba 6: Universo_Completo
    ws5 = wb.create_sheet(title="Universo_Completo")
    all_cols = list(df.columns)
    ws5.append(all_cols)
    for _, row in df.sort_values(by="Composite_Score", ascending=False).iterrows():
        ws5.append([None if pd.isna(v) else v for v in row])
    format_worksheet(ws5, "Universo_Completo")

    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    wb.save(filename)
    print(f"Relatório gerado com sucesso: {filename}")

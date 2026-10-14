"""Exportação para Excel (openpyxl), com a formatação institucional herdada do antigo exporter.

Folhas: ``Resumo_Funil``, ``Top_N`` (ranking, pilares, métricas principais e
explicação), ``Eliminadas`` (ticker, etapa, regra, valor, limiar) e
``Universo_Completo`` (todas as métricas). A Fase 4 acrescenta ``Noticias_IA``.
"""
from __future__ import annotations

import json
import logging
import math
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.funnel.engine import FunnelResult

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")  # azul institucional
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
SUBHEADER_FONT = Font(name="Calibri", size=11, bold=True, color="1F4E79")
DATA_FONT = Font(name="Calibri", size=10)
THIN = Side(border_style="thin", color="E0E0E0")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MIN_COLUMN_WIDTH = 12
MAX_COLUMN_WIDTH = 60
WRAP_COLUMN_WIDTH = 90
WIDTH_SAMPLE_ROWS = 500

# Formatos numéricos por palavra-chave do nome da coluna. A ordem de verificação importa.
PERCENTILE_KEYWORDS = ("_pct", "score", "_z")
PERCENT_KEYWORDS = ("margin", "yield", "roic", "roe", "momentum", "growth", "tax_rate",
                    "price_vs_sma", "cagr", "dilution", "conversion", "equity_to_assets")
RATIO_KEYWORDS = ("_pe", "peg", "ratio", "coverage", "to_ebitda", "to_book", "eps", "altman",
                  "beneish", "ev_to_ebit", "p_ffo")
COUNT_KEYWORDS = ("years", "shares", "days", "piotroski", "count", "stage_order", "entraram",
                  "passaram", "eliminadas", "stages_passed", "rank", "_metrics", "declines",
                  "sector_size")
PRICE_KEYWORDS = ("price", "sma_")
MONEY_KEYWORDS = ("market_cap", "enterprise_value", "revenue", "income", "ebit", "cash", "debt",
                  "assets", "liabilities", "equity", "capex", "profit", "provision", "volume",
                  "working_capital", "receivables", "inventory", "ppe", "interest_expense",
                  "interest_paid", "dividends_paid", "buybacks", "retained", "depreciation",
                  "cost_of", "sga", "ffo")
CENTERED_TEXT_KEYWORDS = ("ticker", "business_type", "sector", "stage", "status", "rule", "source",
                          "currency", "exchange", "country")
WRAP_TEXT_KEYWORDS = ("explanation", "warnings", "all_failed_rules", "description", "classification_note",
                      "ai_red_flags", "ai_catalysts", "ai_risks", "ai_moat", "ai_thesis", "ai_news_impact",
                      "title", "summary")

TOP_SHEET_COLUMNS = (
    "rank", "ticker", "company", "sector", "industry", "business_type", "composite_score",
    "score_quality", "score_valuation", "score_financial_health", "score_momentum", "price",
    "market_cap", "roic_avg_3y", "roe_avg_3y", "fcf_yield", "ev_to_ebit", "forward_pe", "peg_ratio",
    "price_to_book", "dividend_yield", "revenue_cagr_3y", "gross_margin", "operating_margin",
    "cash_conversion", "net_debt_to_ebitda", "interest_coverage", "altman_z", "piotroski_f",
    "beneish_m", "momentum_12_1", "price_vs_sma200", "funnel_warnings", "explanation",
    "news_count", "ai_sentiment", "ai_sentiment_score", "ai_review", "ai_red_flags",
    "ai_thesis_summary", "ai_model",
)
NEWS_SHEET_COLUMNS = ("ticker", "company", "kind", "date", "source", "title", "summary", "link")


def number_format_for(header: str) -> str | None:
    """Formato Excel adequado ao nome (snake_case) da coluna."""
    name = header.lower()
    if "sentiment_score" in name:
        return "0.00"
    if any(k in name for k in PERCENTILE_KEYWORDS):
        return "0.0"
    if any(k in name for k in PERCENT_KEYWORDS):
        return "0.0%"
    if any(k in name for k in RATIO_KEYWORDS):
        return "0.00"
    if any(k in name for k in COUNT_KEYWORDS):
        return "#,##0"
    if any(k in name for k in PRICE_KEYWORDS):
        return "$#,##0.00"
    if any(k in name for k in MONEY_KEYWORDS):
        return '$#,##0,,"M"'
    return None


def excel_value(value: Any) -> Any:
    """Converte um valor pandas/numpy num tipo que o openpyxl aceita (NaN -> vazio)."""
    if value is None:
        return None
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if isinstance(value, (int, np.integer)):
        return int(value)
    if isinstance(value, (float, np.floating)):
        number = float(value)
        return None if math.isnan(number) or math.isinf(number) else number
    if isinstance(value, str):
        return value
    if isinstance(value, (pd.Timestamp, datetime)):
        stamp = pd.Timestamp(value)
        if stamp.tzinfo is not None:
            stamp = stamp.tz_localize(None)
        return stamp.to_pydatetime()
    if isinstance(value, (list, tuple, dict)):
        return json.dumps(value, ensure_ascii=False, default=str)
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    return str(value)


def format_worksheet(ws: Worksheet, header_row: int = 1) -> None:
    """Cabeçalho azul, bordas finas, formatos numéricos por coluna e larguras ajustadas."""
    headers = [str(cell.value or "") for cell in ws[header_row]]
    formats = [number_format_for(h) for h in headers]
    centered = [any(k in h.lower() for k in CENTERED_TEXT_KEYWORDS) for h in headers]
    wrapped = [any(k in h.lower() for k in WRAP_TEXT_KEYWORDS) for h in headers]

    for cell in ws[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 28

    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        for cell in row:
            cell.font = DATA_FONT
            cell.border = CELL_BORDER
            idx = cell.column - 1
            value = cell.value
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                if idx < len(formats) and formats[idx]:
                    cell.number_format = formats[idx]
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                horizontal = "center" if idx < len(centered) and centered[idx] else "left"
                wrap = idx < len(wrapped) and wrapped[idx]
                cell.alignment = Alignment(horizontal=horizontal, vertical="top" if wrap else "center", wrap_text=wrap)

    for col_idx, header in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        if col_idx - 1 < len(wrapped) and wrapped[col_idx - 1]:
            ws.column_dimensions[letter].width = WRAP_COLUMN_WIDTH
            continue
        max_len = len(header)
        last_row = min(ws.max_row, header_row + WIDTH_SAMPLE_ROWS)
        for (cell,) in ws.iter_rows(min_row=header_row + 1, max_row=last_row, min_col=col_idx, max_col=col_idx):
            if cell.value is None:
                continue
            length = 14 if isinstance(cell.value, float) else len(str(cell.value))
            max_len = max(max_len, length)
        ws.column_dimensions[letter].width = min(max(max_len + 3, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if ws.max_row > header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(max(1, ws.max_column))}{ws.max_row}"


def write_dataframe(ws: Worksheet, df: pd.DataFrame) -> None:
    """Escreve um DataFrame (cabeçalho + linhas) e aplica a formatação."""
    ws.append([str(col) for col in df.columns])
    for values in df.itertuples(index=False, name=None):
        ws.append([excel_value(v) for v in values])
    format_worksheet(ws)


def write_funnel_summary(ws: Worksheet, result: FunnelResult, top_rules: int = 10) -> None:
    """Folha Resumo_Funil: contagens por etapa, regras que mais eliminaram e a linha-resumo."""
    stage_table = result.stage_table()
    ws.append([str(col) for col in stage_table.columns])
    for values in stage_table.itertuples(index=False, name=None):
        ws.append([excel_value(v) for v in values])
    format_worksheet(ws)

    ws.append([])
    ws.append(["regra que mais eliminou", "empresas"])
    for cell in ws[ws.max_row]:
        cell.font = SUBHEADER_FONT
    rules = result.top_rules(top_rules)
    for rule, count in rules:
        ws.append([rule, count])
    if not rules:
        ws.append(["(sem eliminações por regra)", 0])

    ws.append([])
    ws.append(["resumo"])
    ws.cell(row=ws.max_row, column=1).font = SUBHEADER_FONT
    for line in result.summary_text().splitlines():
        ws.append([line])
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 70)


def top_sheet_frame(top_df: pd.DataFrame) -> pd.DataFrame:
    columns = [col for col in TOP_SHEET_COLUMNS if col in top_df.columns]
    return top_df[columns].copy()


def news_sheet_frame(top_df: pd.DataFrame, news_by_ticker: dict[str, list] | None, ai_results: dict | None) -> pd.DataFrame:
    """Folha Noticias_IA: por empresa do Top N, uma linha com a análise da IA e uma linha por notícia."""
    rows: list[dict[str, Any]] = []
    names = dict(zip(top_df.get("ticker", []), top_df.get("company", []))) if not top_df.empty else {}
    tickers = list(top_df["ticker"]) if "ticker" in top_df.columns else []
    for ticker in tickers:
        company = names.get(ticker, "")
        result = (ai_results or {}).get(ticker)
        if result is not None:
            analysis = result.analysis or {}
            if result.ok:
                review = f" | {'REVER MANUALMENTE'}" if result.review_required else ""
                title = f"Sentimento: {analysis.get('sentiment')} ({analysis.get('sentiment_score'):+.2f}){review}"
                summary = "\n".join(
                    part for part in (
                        f"Tese: {analysis.get('thesis_summary', '')}",
                        f"Moat: {analysis.get('moat_assessment', '')}",
                        "Catalisadores: " + (" | ".join(analysis.get("catalysts") or []) or "—"),
                        "Riscos: " + (" | ".join(analysis.get("risks") or []) or "—"),
                        "Red flags: " + (" | ".join(analysis.get("red_flags") or []) or "nenhuma"),
                        f"Impacto das notícias: {analysis.get('news_impact', '')}",
                    )
                )
            else:
                title, summary = "Análise IA falhou", str(result.error or "")
            rows.append({"ticker": ticker, "company": company, "kind": "ANÁLISE IA", "date": result.analyzed_at[:10],
                         "source": result.model, "title": title, "summary": summary, "link": ""})
        for item in (news_by_ticker or {}).get(ticker, []):
            rows.append({"ticker": ticker, "company": company, "kind": "NOTÍCIA", "date": item.published_text,
                         "source": item.source, "title": item.title, "summary": item.summary, "link": item.link})
    return pd.DataFrame(rows, columns=list(NEWS_SHEET_COLUMNS))


def save_workbook(wb: openpyxl.Workbook, path: Path) -> Path:
    """Grava o livro; se o ficheiro estiver aberto no Excel, grava uma cópia com hora no nome."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
        return path
    except PermissionError:
        alternative = path.with_name(f"{path.stem}_{datetime.now():%H%M%S}{path.suffix}")
        log.warning("Sem permissão para escrever %s (ficheiro aberto?); a gravar em %s", path.name, alternative.name)
        wb.save(alternative)
        return alternative


def export_screener_workbook(
    result: FunnelResult,
    metrics_df: pd.DataFrame,
    path: Path,
    top_df: pd.DataFrame | None = None,
    news_by_ticker: dict[str, list] | None = None,
    ai_results: dict | None = None,
) -> Path:
    """Livro Excel: Resumo_Funil, Top_N, Noticias_IA (se houver notícias/IA), Eliminadas e Universo_Completo."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    write_funnel_summary(wb.create_sheet("Resumo_Funil"), result)
    if top_df is not None:
        write_dataframe(wb.create_sheet("Top_N"), top_sheet_frame(top_df))
        if news_by_ticker or ai_results:
            write_dataframe(wb.create_sheet("Noticias_IA"), news_sheet_frame(top_df, news_by_ticker, ai_results))
    write_dataframe(wb.create_sheet("Eliminadas"), result.eliminations_table())
    write_dataframe(wb.create_sheet("Universo_Completo"), metrics_df)

    saved = save_workbook(wb, path)
    log.info("Excel gravado: %s (%d folhas)", saved, len(wb.sheetnames))
    return saved
